import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.comments import Comment
from typing import Dict, Optional
from datetime import datetime
from io import BytesIO
from src.utils.filters import generate_filter_description, generate_filter_metadata
from config import setup_logger

logger = setup_logger()


class ExcelExporter:
    def __init__(self):
        self.header_font = Font(name="PingFang SC", size=11, bold=True, color="FFFFFFFF")
        self.header_fill = PatternFill(start_color="FF1E3A8A", end_color="FF1E3A8A", fill_type="solid")
        self.watermark_font = Font(name="PingFang SC", size=9, color="FF6B7280", italic=True)
        self.title_font = Font(name="PingFang SC", size=14, bold=True, color="FF1E3A8A")
        self.thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

    def export_to_excel(
        self,
        data: pl.DataFrame,
        filters: Optional[Dict] = None,
        sheet_name: str = "数据导出",
        title: str = "教材订购数据导出"
    ) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        self._add_watermark(ws, filters)
        self._add_title(ws, title, filters)

        start_row = 6
        pandas_df = data.to_pandas()

        for r_idx, row in enumerate(dataframe_to_rows(pandas_df, index=False, header=True), start=start_row):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.border = self.thin_border

                if r_idx == start_row:
                    cell.font = self.header_font
                    cell.fill = self.header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max(max_length + 2, 10), 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        logger.info(f"成功生成Excel文件，包含{len(data)}条记录")
        return output.getvalue()

    def _add_watermark(self, ws, filters: Optional[Dict]):
        watermark_text = self._generate_watermark_text(filters)
        comment = Comment(watermark_text, "系统导出")
        comment.width = 400
        comment.height = 150
        ws.cell(row=1, column=1).comment = comment

    def _add_title(self, ws, title: str, filters: Optional[Dict]):
        ws.cell(row=1, column=1, value=title).font = self.title_font
        ws.merge_cells("A1:Z1")

        filter_desc = generate_filter_description(filters)
        ws.cell(row=2, column=1, value=f"筛选条件: {filter_desc}").font = self.watermark_font
        ws.merge_cells("A2:Z2")

        export_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.cell(row=3, column=1, value=f"导出时间: {export_time}").font = self.watermark_font
        ws.merge_cells("A3:Z3")

        metadata = generate_filter_metadata(filters)
        ws.cell(row=4, column=1, value=f"取数范围标识: {self._generate_identifier(metadata)}").font = self.watermark_font
        ws.merge_cells("A4:Z4")

    def _generate_watermark_text(self, filters: Optional[Dict]) -> str:
        desc = generate_filter_description(filters)
        time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        return f"""
        高校教务教材订购漏斗报表系统
        -----------------------------
        导出时间: {time}
        筛选条件: {desc}
        本文件数据基于上述筛选条件导出，
        转发时请保留此信息以确保数据可追溯。
        -----------------------------
        """

    def _generate_identifier(self, metadata: Dict) -> str:
        import hashlib
        import json
        filter_str = json.dumps(metadata["filters"], sort_keys=True, ensure_ascii=False)
        hash_val = hashlib.md5(filter_str.encode("utf-8")).hexdigest()[:8]
        return f"{datetime.now().strftime('%Y%m%d')}_{hash_val}"

    def export_multiple_sheets(
        self,
        sheets_data: Dict[str, pl.DataFrame],
        filters: Optional[Dict] = None,
        title: str = "教材订购数据分析报告"
    ) -> bytes:
        wb = Workbook()
        wb.remove(wb.active)

        for sheet_name, data in sheets_data.items():
            ws = wb.create_sheet(title=sheet_name)
            self._add_title(ws, f"{title} - {sheet_name}", filters)

            start_row = 6
            pandas_df = data.to_pandas()

            for r_idx, row in enumerate(dataframe_to_rows(pandas_df, index=False, header=True), start=start_row):
                for c_idx, value in enumerate(row, start=1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=value)
                    cell.border = self.thin_border

                    if r_idx == start_row:
                        cell.font = self.header_font
                        cell.fill = self.header_fill
                        cell.alignment = Alignment(horizontal="center", vertical="center")

            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max(max_length + 2, 10), 50)
                ws.column_dimensions[column_letter].width = adjusted_width

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        logger.info(f"成功生成多Sheet Excel文件，包含{len(sheets_data)}个工作表")
        return output.getvalue()


excel_exporter = ExcelExporter()

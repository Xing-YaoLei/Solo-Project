import os
import logging
from datetime import datetime
from typing import Optional, Dict, Any, List
import pandas as pd
import numpy as np

import xlsxwriter
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import Config
from app.services.data_service import (
    QueryService, TrendService, FinanceService, RefreshService
)
from app.services.auth_service import AuthService
from app.models import User

logger = logging.getLogger(__name__)


class ExportService:
    def __init__(self, config: Optional[Config] = None):
        self.config = config or Config
        self.export_dir = self.config.EXPORT_DIR
        os.makedirs(self.export_dir, exist_ok=True)

    def export_dashboard_data(
        self,
        user: User,
        filters: Optional[Dict[str, Any]] = None,
        include_finance: Optional[bool] = None,
        format: str = 'xlsx'
    ) -> str:
        config_dict = {
            'ROLE_PERMISSIONS': self.config.ROLE_PERMISSIONS,
            'EXPORT_CALIBER_NOTES': self.config.EXPORT_CALIBER_NOTES
        }
        if include_finance is None:
            include_finance = AuthService.can_view_finance(user)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        role_suffix = user.role
        filename = f"法律案件风险监测数据_{role_suffix}_{timestamp}.{format}"
        filepath = os.path.join(self.export_dir, filename)

        try:
            cases_df = QueryService.get_cases_df(user, config_dict, include_finance=include_finance)
            clients_df = QueryService.get_clients_df(user, config_dict)
            evidences_df = QueryService.get_evidences_df(user, config_dict)
            hearings_df = QueryService.get_hearings_df(user, config_dict)
            payments_df = QueryService.get_payments_df(user, config_dict, include_finance=include_finance)

            client_trend_df = TrendService.get_client_trend_df(clients_df)
            case_stage_df = TrendService.get_case_stage_distribution(cases_df)
            evidence_detail_df = TrendService.get_evidence_detail_df(evidences_df, cases_df)
            hearing_anomaly_df = TrendService.get_hearing_anomaly_df(hearings_df, cases_df)
            payment_summary_df = FinanceService.get_payment_summary_df(payments_df, cases_df)

            if filters:
                cases_df = self._apply_filters(cases_df, filters)
                hearings_df = self._apply_filters(hearings_df, filters)
                evidences_df = self._apply_filters(evidences_df, filters)
                payments_df = self._apply_filters(payments_df, filters)
                evidence_detail_df = self._apply_filters(evidence_detail_df, filters)
                hearing_anomaly_df = self._apply_filters(hearing_anomaly_df, filters)
                payment_summary_df = self._apply_filters(payment_summary_df, filters)

            if format == 'xlsx':
                self._write_xlsx(
                    filepath, user, include_finance,
                    cases_df, clients_df, evidences_df, hearings_df, payments_df,
                    client_trend_df, case_stage_df, evidence_detail_df,
                    hearing_anomaly_df, payment_summary_df
                )
            else:
                self._write_csvs(filepath, cases_df, hearings_df, evidences_df, payments_df)

            return filepath
        except Exception as e:
            logger.error(f"Export error: {e}")
            raise

    def _apply_filters(self, df: pd.DataFrame, filters: Dict[str, Any]) -> pd.DataFrame:
        if df.empty:
            return df
        result = df.copy()
        for col, values in filters.items():
            if col in result.columns and values:
                if isinstance(values, list):
                    if 'id' in col:
                        result = result[result[col].isin(values)]
                    else:
                        result = result[result[col].astype(str).isin([str(v) for v in values])]
                else:
                    result = result[result[col].astype(str) == str(values)]
        return result

    def _write_xlsx(
        self, filepath: str, user: User, include_finance: bool,
        cases_df, clients_df, evidences_df, hearings_df, payments_df,
        client_trend_df, case_stage_df, evidence_detail_df,
        hearing_anomaly_df, payment_summary_df
    ):
        writer = pd.ExcelWriter(filepath, engine='xlsxwriter')
        workbook = writer.book

        header_format = workbook.add_format({
            'bold': True,
            'bg_color': '#1F4E79',
            'font_color': 'white',
            'align': 'center',
            'valign': 'vcenter',
            'border': 1,
            'font_size': 11
        })
        caliber_format = workbook.add_format({
            'italic': True,
            'font_color': '#808080',
            'font_size': 9,
            'text_wrap': True
        })
        title_format = workbook.add_format({
            'bold': True,
            'font_size': 14,
            'font_color': '#1F4E79'
        })
        subheader_format = workbook.add_format({
            'bold': True,
            'bg_color': '#D6E4F0',
            'font_size': 10,
            'border': 1
        })

        self._write_caliber_sheet(writer, workbook, title_format, caliber_format, subheader_format)

        self._write_sheet_with_title(
            writer, '客户档案趋势', client_trend_df,
            title_format, header_format,
            '新增客户数量趋势（按月统计）'
        )

        self._write_sheet_with_title(
            writer, '案件阶段构成', case_stage_df,
            title_format, header_format,
            '当前案件在各阶段的分布情况'
        )

        if not evidence_detail_df.empty:
            evidence_display_cols = [
                'evidence_code', 'evidence_name', 'evidence_type', 'evidence_category',
                'status', 'case_number', 'case_name', 'file_name', 'file_size',
                'page_count', 'submitted_to_court', 'authenticity_score',
                'relevance_score', 'created_at'
            ]
            evidence_display = evidence_detail_df[
                [c for c in evidence_display_cols if c in evidence_detail_df.columns]
            ]
            evidence_display = evidence_display.rename(columns={
                'evidence_code': '证据编号',
                'evidence_name': '证据名称',
                'evidence_type': '证据类型',
                'evidence_category': '证据分类',
                'status': '证据状态',
                'case_number': '案件编号',
                'case_name': '案件名称',
                'file_name': '文件名',
                'file_size': '文件大小(字节)',
                'page_count': '页数',
                'submitted_to_court': '是否已提交',
                'authenticity_score': '真实性评分',
                'relevance_score': '关联性评分',
                'created_at': '入库时间'
            })
        else:
            evidence_display = pd.DataFrame(columns=['暂无证据数据'])

        self._write_sheet_with_title(
            writer, '证据附件明细', evidence_display,
            title_format, header_format,
            '案件相关证据材料清单及状态追踪'
        )

        if not hearing_anomaly_df.empty:
            hearing_display_cols = [
                '异常类型', '紧急程度', 'case_number', 'case_name',
                'current_stage', 'hearing_type', 'scheduled_at',
                'location', 'status', 'preparation_status', '主办律师'
            ]
            hearing_display = hearing_anomaly_df[
                [c for c in hearing_display_cols if c in hearing_anomaly_df.columns]
            ]
            hearing_display = hearing_display.rename(columns={
                'case_number': '案件编号',
                'case_name': '案件名称',
                'current_stage': '案件阶段',
                'hearing_type': '庭审类型',
                'scheduled_at': '开庭时间',
                'location': '开庭地点',
                'status': '庭审状态',
                'preparation_status': '准备进度',
            })
        else:
            hearing_display = pd.DataFrame(columns=['暂无庭审异常记录'])

        self._write_sheet_with_title(
            writer, '庭审日程异常', hearing_display,
            title_format, header_format,
            '存在异常情况的庭审日程清单（含风险标注）'
        )

        case_display = self._prepare_cases_for_export(cases_df, include_finance)
        self._write_sheet_with_title(
            writer, '案件总览', case_display,
            title_format, header_format,
            '委托案件基础信息与当前状态'
        )

        if include_finance and not payment_summary_df.empty:
            fin_display = payment_summary_df.rename(columns={
                'case_number': '案件编号',
                'case_name': '案件名称',
                'client_name': '客户名称',
                'current_stage': '案件阶段',
                '合同总额': '合同金额(元)',
                '应收总额': '应收金额(元)',
                '实收总额': '实收金额(元)',
                '笔数': '收款笔数',
                '回款率': '回款率(%)',
                '差额': '未收差额(元)'
            })
            self._write_sheet_with_title(
                writer, '收款汇总', fin_display,
                title_format, header_format,
                '各案件委托合同款项收取进度明细'
            )

        self._write_info_sheet(
            writer, workbook, title_format, caliber_format,
            user, include_finance
        )

        writer.close()
        self._post_format_xlsx(filepath)

    def _write_caliber_sheet(self, writer, workbook, title_fmt, caliber_fmt, sub_fmt):
        sheet_name = '数据口径说明'
        ws = workbook.add_worksheet(sheet_name)
        row = 0
        ws.set_column('A:A', 18)
        ws.set_column('B:B', 80)

        ws.write(row, 0, '法律服务案件委托数据口径说明书', title_fmt)
        ws.merge_range(row, 0, row, 1, '', title_fmt)
        row += 2

        caliber_notes = self.config.EXPORT_CALIBER_NOTES

        for section, items in caliber_notes.items():
            ws.write(row, 0, section, sub_fmt)
            ws.write(row, 1, section, sub_fmt)
            ws.merge_range(row, 0, row, 1, section, sub_fmt)
            row += 1

            if isinstance(items, dict):
                for key, desc in items.items():
                    ws.write(row, 0, key)
                    ws.write(row, 1, desc, caliber_fmt)
                    row += 1
            row += 1

        ws.write(row, 0, '数据来源', sub_fmt)
        ws.merge_range(row, 0, row, 1, '数据来源', sub_fmt)
        row += 1
        sources = [
            ('客户档案', '律所客户管理系统（CMS）建档数据'),
            ('案件委托', '委托合同签订信息及案件进度更新'),
            ('证据附件', '邮件附件自动识别 + 人工录入归档'),
            ('庭审日程', '日历系统（ICS/WebDAV）同步 + 人工维护'),
            ('收款流水', '银行对账单/支付宝/微信账单导入自动匹配'),
        ]
        for src, desc in sources:
            ws.write(row, 0, src)
            ws.write(row, 1, desc, caliber_fmt)
            row += 1

        row += 1
        ws.write(row, 0, '更新时间', sub_fmt)
        ws.merge_range(row, 0, row, 1, '更新时间', sub_fmt)
        row += 1
        refresh_times = RefreshService.get_last_refresh_time()
        labels = [
            ('overall', '总刷新时间'),
            ('email', '邮件附件同步'),
            ('calendar', '日历同步'),
            ('payment', '收款流水导入'),
            ('mv', '物化视图刷新'),
        ]
        for key, label in labels:
            val = refresh_times.get(key) or '暂无记录'
            ws.write(row, 0, label)
            ws.write(row, 1, str(val), caliber_fmt)
            row += 1

        writer.sheets[sheet_name] = ws

    def _write_sheet_with_title(self, writer, sheet_name: str, df: pd.DataFrame,
                                 title_fmt, header_fmt, description: str = ''):
        max_sheet_len = 31
        safe_name = sheet_name[:max_sheet_len]
        start_row = 3 if description else 1

        df.to_excel(writer, sheet_name=safe_name, startrow=start_row, index=False)
        workbook = writer.book
        ws = writer.sheets[safe_name]

        ws.write(0, 0, safe_name, title_fmt)
        if description:
            ws.set_row(1, None, None)
            ws.write(1, 0, description)

        if not df.empty:
            for col_idx in range(len(df.columns)):
                ws.write(start_row, col_idx, df.columns[col_idx], header_fmt)
                max_len = max(
                    df[df.columns[col_idx]].astype(str).map(len).max() if len(df) > 0 else 0,
                    len(str(df.columns[col_idx]))
                )
                ws.set_column(col_idx, col_idx, min(max_len + 2, 50))

        ws.freeze_panes(start_row + 1, 0)

    def _prepare_cases_for_export(self, cases_df: pd.DataFrame, include_finance: bool) -> pd.DataFrame:
        if cases_df.empty:
            return pd.DataFrame(columns=['暂无案件数据'])
        cols = [
            'case_number', 'case_name', 'case_type', 'case_category',
            'client_name', 'opposing_party', 'responsible_lawyer_name',
            'entrusted_at', 'accepted_at', 'filed_at', 'current_stage',
            'stage_updated_at', 'court', 'risk_assessment', 'status'
        ]
        if include_finance:
            cols.extend(['claim_amount', 'judged_amount'])
        display = cases_df[[c for c in cols if c in cases_df.columns]].copy()
        rename_map = {
            'case_number': '案件编号',
            'case_name': '案件名称',
            'case_type': '案件类型',
            'case_category': '案件分类',
            'client_name': '委托人',
            'opposing_party': '对方当事人',
            'responsible_lawyer_name': '主办律师',
            'entrusted_at': '委托日期',
            'accepted_at': '受理日期',
            'filed_at': '立案日期',
            'current_stage': '当前阶段',
            'stage_updated_at': '阶段更新时间',
            'court': '受理法院',
            'risk_assessment': '风险等级',
            'status': '案件状态',
            'claim_amount': '诉讼标的(元)',
            'judged_amount': '判决金额(元)'
        }
        display = display.rename(columns=rename_map)
        return display

    def _write_info_sheet(self, writer, workbook, title_fmt, caliber_fmt,
                          user: User, include_finance: bool):
        sheet_name = '导出信息'
        ws = workbook.add_worksheet(sheet_name)
        ws.set_column('A:A', 20)
        ws.set_column('B:B', 60)

        row = 0
        ws.write(row, 0, '导出信息记录', title_fmt)
        ws.merge_range(row, 0, row, 1, '', title_fmt)
        row += 2

        info_items = [
            ('导出时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
            ('导出人', user.full_name),
            ('导出人角色', user.role),
            ('导出人账号', user.username),
            ('数据范围', self._get_scope_description(user)),
            ('是否含财务数据', '是' if include_finance else '否'),
            ('数据口径版本', 'v1.0'),
            ('生成系统', 'Legal Risk Monitor Dashboard'),
            ('保密提示', '本文件含客户敏感信息，仅限授权人员查阅，不得外传'),
        ]
        for label, value in info_items:
            ws.write(row, 0, label)
            ws.write(row, 1, str(value), caliber_fmt)
            row += 1

        writer.sheets[sheet_name] = ws

    def _get_scope_description(self, user: User) -> str:
        role = user.role
        scope_map = {
            'admin': '全所全部案件数据',
            'partner': f'所在部门({user.department or "无"})全部案件',
            'lawyer': '本人主办/负责的案件',
            'paralegal': '所在团队参与的案件',
            'client': '本人/本公司委托的案件',
            'auditor': '全所只读数据（审计用）',
        }
        return scope_map.get(role, '默认范围')

    def _post_format_xlsx(self, filepath: str):
        try:
            wb = load_workbook(filepath)
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            if not cell.border or cell.border == Border():
                                cell.border = thin_border
                            if not cell.alignment or cell.alignment == Alignment():
                                cell.alignment = Alignment(vertical='center')
            wb.save(filepath)
        except Exception as e:
            logger.warning(f"Post-format xlsx skipped: {e}")

    def _write_csvs(self, basepath: str, cases_df, hearings_df, evidences_df, payments_df):
        base_dir = os.path.splitext(basepath)[0]
        os.makedirs(base_dir, exist_ok=True)
        cases_df.to_csv(os.path.join(base_dir, '案件总览.csv'), index=False, encoding='utf-8-sig')
        hearings_df.to_csv(os.path.join(base_dir, '庭审日程.csv'), index=False, encoding='utf-8-sig')
        evidences_df.to_csv(os.path.join(base_dir, '证据明细.csv'), index=False, encoding='utf-8-sig')
        payments_df.to_csv(os.path.join(base_dir, '收款流水.csv'), index=False, encoding='utf-8-sig')

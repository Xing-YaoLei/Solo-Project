import os
from datetime import datetime

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from api.models import ExportRecord, Prescription


async def create_export_record(db: AsyncSession, task_id: str, export_type: str, user_id: str) -> ExportRecord:
    record = ExportRecord(
        task_id=task_id,
        export_type=export_type,
        requested_by=user_id,
    )
    db.add(record)
    await db.flush()
    return record


async def get_export_record_by_task(db: AsyncSession, task_id: str) -> ExportRecord | None:
    result = await db.execute(select(ExportRecord).where(ExportRecord.task_id == task_id))
    return result.scalar_one_or_none()


async def update_export_record(
    db: AsyncSession,
    record: ExportRecord,
    *,
    status: str | None = None,
    file_path: str | None = None,
    error_message: str | None = None,
    row_count: int | None = None,
) -> ExportRecord:
    if status is not None:
        record.status = status
    if file_path is not None:
        record.file_path = file_path
    if error_message is not None:
        record.error_message = error_message
    if row_count is not None:
        record.row_count = row_count
    if status in ("completed", "failed"):
        record.completed_at = datetime.utcnow()
    await db.flush()
    return record


def generate_excel(
    prescriptions: list[dict],
    caliber_notes: list[dict],
    export_dir: str,
    task_id: str,
) -> tuple[str, int]:
    os.makedirs(export_dir, exist_ok=True)

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    ws_prescriptions = wb.active
    ws_prescriptions.title = "处方数据"

    prescription_headers = [
        "处方编号", "状态", "处方类型", "门店名称", "患者姓名",
        "诊断", "总金额", "提交时间", "审核时间", "审核人",
        "批号信息", "医保结算状态", "补货状态",
    ]
    ws_prescriptions.append(prescription_headers)

    header_font = Font(bold=True)
    for col in range(1, len(prescription_headers) + 1):
        cell = ws_prescriptions.cell(row=1, column=col)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    column_widths = [18, 12, 12, 22, 12, 20, 10, 20, 20, 12, 30, 14, 12]
    for i, width in enumerate(column_widths, 1):
        ws_prescriptions.column_dimensions[get_column_letter(i)].width = width

    ws_prescriptions.freeze_panes = "A2"

    status_map = {
        "pending": "待审核",
        "in_review": "审核中",
        "approved": "已通过",
        "rejected": "已驳回",
        "exception": "异常",
    }

    row_count = 0
    for p in prescriptions:
        batch_info = p.get("batch_items", [])
        batch_json = str(batch_info) if batch_info else ""

        insurance_status = "未结算"
        if p.get("insurance_records"):
            for ins in p["insurance_records"]:
                if ins.get("verified"):
                    insurance_status = "已结算"
                    break

        replenishment_status = "无"
        if p.get("replenishments"):
            statuses = [r.get("status", "") for r in p["replenishments"]]
            if "pending" in statuses:
                replenishment_status = "补货中"
            elif "fulfilled" in statuses:
                replenishment_status = "已补货"

        ws_prescriptions.append([
            p.get("rx_number", ""),
            status_map.get(p.get("status", ""), p.get("status", "")),
            p.get("priority", ""),
            p.get("store_name", ""),
            p.get("patient_name", ""),
            p.get("diagnosis", ""),
            p.get("total_amount", 0),
            str(p.get("created_at", "")),
            str(p.get("reviewed_at", "")),
            p.get("reviewer_name", ""),
            batch_json,
            insurance_status,
            replenishment_status,
        ])
        row_count += 1

    ws_caliber = wb.create_sheet(title="口径说明")

    caliber_headers = ["指标名称", "定义说明", "排除项", "备注"]
    ws_caliber.append(caliber_headers)

    for col in range(1, len(caliber_headers) + 1):
        cell = ws_caliber.cell(row=1, column=col)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    caliber_widths = [20, 50, 40, 30]
    for i, width in enumerate(caliber_widths, 1):
        ws_caliber.column_dimensions[get_column_letter(i)].width = width

    ws_caliber.freeze_panes = "A2"

    for note in caliber_notes:
        exclusions = note.get("exclusions", [])
        exclusions_str = "\n".join(f"- {e}" for e in exclusions) if exclusions else ""
        ws_caliber.append([
            note.get("metric", ""),
            note.get("definition", ""),
            exclusions_str,
            note.get("remarks", ""),
        ])

    filename = f"export_{task_id[:8]}_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}.xlsx"
    file_path = os.path.join(export_dir, filename)
    wb.save(file_path)
    return file_path, row_count


async def update_export_record_async(
    db: AsyncSession,
    task_id: str,
    *,
    status: str | None = None,
    file_path: str | None = None,
    error_message: str | None = None,
    row_count: int | None = None,
) -> ExportRecord | None:
    result = await db.execute(select(ExportRecord).where(ExportRecord.task_id == task_id))
    record = result.scalar_one_or_none()
    if not record:
        return None
    return await update_export_record(
        db,
        record,
        status=status,
        file_path=file_path,
        error_message=error_message,
        row_count=row_count,
    )

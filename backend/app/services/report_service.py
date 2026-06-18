from sqlalchemy.orm import Session
from datetime import datetime, date
from dateutil.relativedelta import relativedelta
from typing import Optional, Dict, Any, List
from ..models import WorkOrder, WorkOrderStatus, ReworkRecord, ReportDownloadLog, User, Vehicle
from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


def calculate_rework_rate(
    db: Session,
    year: int,
    month: int,
    technician_id: Optional[int] = None,
    station_id: Optional[int] = None,
) -> Dict[str, Any]:
    start_date = date(year, month, 1)
    end_date = start_date + relativedelta(months=1)

    query = db.query(WorkOrder).filter(
        WorkOrder.created_at >= datetime.combine(start_date, datetime.min.time()),
        WorkOrder.created_at < datetime.combine(end_date, datetime.min.time()),
    )

    if technician_id:
        query = query.filter(WorkOrder.technician_id == technician_id)
    if station_id:
        query = query.filter(WorkOrder.station_id == station_id)

    all_orders = query.all()
    total_orders = len(all_orders)
    rework_orders = [o for o in all_orders if o.is_rework]
    rework_count = len(rework_orders)
    rework_rate = (rework_count / total_orders * 100) if total_orders > 0 else 0.0

    details = []
    for order in rework_orders:
        vehicle = db.query(Vehicle).filter(Vehicle.id == order.vehicle_id).first()
        rework_record = db.query(ReworkRecord).filter(ReworkRecord.rework_order_id == order.id).first()
        details.append({
            "order_no": order.order_no,
            "plate_number": vehicle.plate_number if vehicle else "",
            "brand_model": f"{vehicle.brand} {vehicle.model}" if vehicle else "",
            "complaint": order.complaint or "",
            "technician": order.technician.full_name if order.technician else "",
            "reason": rework_record.reason if rework_record else "",
            "created_at": order.created_at.strftime("%Y-%m-%d %H:%M"),
        })

    return {
        "month": f"{year}-{month:02d}",
        "total_orders": total_orders,
        "rework_orders": rework_count,
        "rework_rate": round(rework_rate, 2),
        "details": details,
    }


def generate_rework_excel_report(
    db: Session,
    year: int,
    month: int,
    filter_criteria: Dict[str, Any],
    generated_by: int,
) -> BytesIO:
    data = calculate_rework_rate(
        db,
        year,
        month,
        technician_id=filter_criteria.get("technician_id"),
        station_id=filter_criteria.get("station_id"),
    )

    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "返修率报表"

    title_font = Font(bold=True, size=16)
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    ws["A1"] = f"返修率报表 - {data['month']}"
    ws["A1"].font = title_font
    ws.merge_cells("A1:F1")
    ws["A1"].alignment = center_align

    ws["A3"] = "筛选条件"
    ws["A3"].font = Font(bold=True, size=12)
    ws["A4"] = json.dumps(filter_criteria, ensure_ascii=False, indent=2)

    user = db.query(User).filter(User.id == generated_by).first()
    ws["A6"] = f"生成者: {user.full_name if user else 'Unknown'}"
    ws["B6"] = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

    ws["A8"] = "统计汇总"
    ws["A8"].font = Font(bold=True, size=12)
    ws["A9"] = "总工单"
    ws["B9"] = data["total_orders"]
    ws["C9"] = "返修工单"
    ws["D9"] = data["rework_orders"]
    ws["E9"] = "返修率"
    ws["F9"] = f"{data['rework_rate']}%"

    headers = ["工单号", "车牌号", "车型", "故障描述", "技师", "返修原因", "创建时间"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=12, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for row_idx, detail in enumerate(data["details"], 13):
        ws.cell(row=row_idx, column=1, value=detail["order_no"])
        ws.cell(row=row_idx, column=2, value=detail["plate_number"])
        ws.cell(row=row_idx, column=3, value=detail["brand_model"])
        ws.cell(row=row_idx, column=4, value=detail["complaint"])
        ws.cell(row=row_idx, column=5, value=detail["technician"])
        ws.cell(row=row_idx, column=6, value=detail["reason"])
        ws.cell(row=row_idx, column=7, value=detail["created_at"])

    for col in ["A", "B", "C", "D", "E", "F", "G"]:
        ws.column_dimensions[col].width = 20

    wb.save(output)
    output.seek(0)
    return output


def log_report_download(
    db: Session,
    report_type: str,
    filter_criteria: Dict[str, Any],
    generated_by: int,
    file_name: str,
) -> ReportDownloadLog:
    log = ReportDownloadLog(
        report_type=report_type,
        filter_criteria=filter_criteria,
        generated_by=generated_by,
        file_name=file_name,
    )
    db.add(log)
    db.commit()
    db.refresh(log)
    return log


def get_report_history(db: Session, limit: int = 50) -> List[ReportDownloadLog]:
    return (
        db.query(ReportDownloadLog)
        .order_by(ReportDownloadLog.created_at.desc())
        .limit(limit)
        .all()
    )

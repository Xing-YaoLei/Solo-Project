from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django.db.models import Count, Avg, Q
from django.utils import timezone
from datetime import datetime, timedelta
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse
from .models import DownloadRecord, MonthlyReview
from .serializers import DownloadRecordSerializer, MonthlyReviewSerializer


class DownloadRecordViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = DownloadRecord.objects.all()
    serializer_class = DownloadRecordSerializer
    filterset_fields = ['report_type', 'user']
    ordering_fields = ['generated_at']


class MonthlyReviewViewSet(viewsets.ModelViewSet):
    queryset = MonthlyReview.objects.all()
    serializer_class = MonthlyReviewSerializer
    filterset_fields = ['year', 'month', 'counselor']
    ordering_fields = ['year', 'month']

    @action(detail=False, methods=['get'])
    def completion_rate(self, request):
        year = int(request.query_params.get('year', timezone.now().year))
        month = int(request.query_params.get('month', timezone.now().month))
        counselor_id = request.query_params.get('counselor')

        from apps.followups.models import RenewalFollowUp
        from apps.courses.models import Enrollment
        from apps.users.models import User

        first_day = datetime(year, month, 1)
        if month == 12:
            last_day = datetime(year + 1, 1, 1) - timedelta(seconds=1)
        else:
            last_day = datetime(year, month + 1, 1) - timedelta(seconds=1)

        follow_ups = RenewalFollowUp.objects.filter(
            created_at__gte=first_day,
            created_at__lte=last_day
        )

        if counselor_id:
            follow_ups = follow_ups.filter(counselor_id=counselor_id)

        total = follow_ups.count()
        completed = follow_ups.filter(status__in=['completed', 'renewed', 'closed']).count()
        renewed = follow_ups.filter(status='renewed').count()
        completion_rate = (completed / total * 100) if total > 0 else 0
        renewal_rate = (renewed / total * 100) if total > 0 else 0

        enrollments = Enrollment.objects.filter(status='active')
        if counselor_id:
            enrollments = enrollments.filter(student__student_profile__counselor_id=counselor_id)

        avg_progress = enrollments.aggregate(avg=Avg('progress'))['avg'] or 0
        total_students = enrollments.values('student').distinct().count()

        status_breakdown = list(
            follow_ups.values('status').annotate(count=Count('id'))
        )

        priority_breakdown = list(
            follow_ups.values('priority').annotate(count=Count('id'))
        )

        data = {
            'year': year,
            'month': month,
            'total_follow_ups': total,
            'completed_follow_ups': completed,
            'renewal_count': renewed,
            'completion_rate': round(completion_rate, 2),
            'renewal_rate': round(renewal_rate, 2),
            'average_student_progress': round(avg_progress, 2),
            'total_students': total_students,
            'status_breakdown': status_breakdown,
            'priority_breakdown': priority_breakdown,
        }

        return Response(data)

    @action(detail=False, methods=['post'])
    def generate_monthly_review(self, request):
        year = int(request.data.get('year', timezone.now().year))
        month = int(request.data.get('month', timezone.now().month))
        counselor_id = request.data.get('counselor')
        summary = request.data.get('summary', '')

        from apps.followups.models import RenewalFollowUp
        from apps.courses.models import Enrollment

        first_day = datetime(year, month, 1)
        if month == 12:
            last_day = datetime(year + 1, 1, 1) - timedelta(seconds=1)
        else:
            last_day = datetime(year, month + 1, 1) - timedelta(seconds=1)

        follow_ups = RenewalFollowUp.objects.filter(
            created_at__gte=first_day,
            created_at__lte=last_day
        )

        if counselor_id:
            follow_ups = follow_ups.filter(counselor_id=counselor_id)

        total = follow_ups.count()
        completed = follow_ups.filter(status__in=['completed', 'renewed', 'closed']).count()
        renewed = follow_ups.filter(status='renewed').count()
        completion_rate = (completed / total * 100) if total > 0 else 0
        renewal_rate = (renewed / total * 100) if total > 0 else 0

        enrollments = Enrollment.objects.filter(status='active')
        if counselor_id:
            enrollments = enrollments.filter(student__student_profile__counselor_id=counselor_id)

        avg_progress = enrollments.aggregate(avg=Avg('progress'))['avg'] or 0
        total_students = enrollments.values('student').distinct().count()

        review, created = MonthlyReview.objects.update_or_create(
            year=year,
            month=month,
            counselor_id=counselor_id,
            defaults={
                'total_follow_ups': total,
                'completed_follow_ups': completed,
                'renewal_count': renewed,
                'completion_rate': round(completion_rate, 2),
                'renewal_rate': round(renewal_rate, 2),
                'average_student_progress': round(avg_progress, 2),
                'total_students': total_students,
                'summary': summary,
                'generated_by': request.user
            }
        )

        return Response(MonthlyReviewSerializer(review).data)

    @action(detail=False, methods=['post'])
    def export_follow_ups(self, request):
        filters = request.data.get('filters', {})
        report_type = 'follow_ups'

        from apps.followups.models import RenewalFollowUp

        queryset = RenewalFollowUp.objects.all()

        if filters.get('status'):
            queryset = queryset.filter(status=filters['status'])
        if filters.get('priority'):
            queryset = queryset.filter(priority=filters['priority'])
        if filters.get('counselor'):
            queryset = queryset.filter(counselor_id=filters['counselor'])
        if filters.get('start_date'):
            queryset = queryset.filter(created_at__gte=filters['start_date'])
        if filters.get('end_date'):
            queryset = queryset.filter(created_at__lte=filters['end_date'])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '续费跟进记录'

        headers = ['ID', '学员姓名', '咨询师', '课程', '状态', '优先级',
                   '原因', '进度', '续费到期日', '创建时间', '更新时间']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        status_map = dict(RenewalFollowUp.STATUS_CHOICES)
        priority_map = dict(RenewalFollowUp.PRIORITY_CHOICES)

        for row, item in enumerate(queryset, 2):
            ws.cell(row=row, column=1, value=item.id)
            ws.cell(row=row, column=2, value=item.student.username if item.student else '')
            ws.cell(row=row, column=3, value=item.counselor.username if item.counselor else '')
            ws.cell(row=row, column=4,
                    value=item.enrollment.course.name if item.enrollment and item.enrollment.course else '')
            ws.cell(row=row, column=5, value=status_map.get(item.status, item.status))
            ws.cell(row=row, column=6, value=priority_map.get(item.priority, item.priority))
            ws.cell(row=row, column=7, value=item.reason)
            ws.cell(row=row, column=8, value=f'{item.enrollment.progress}%' if item.enrollment else '')
            ws.cell(row=row, column=9,
                    value=item.enrollment.renewal_due_date if item.enrollment else '')
            ws.cell(row=row, column=10, value=item.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=11, value=item.updated_at.strftime('%Y-%m-%d %H:%M:%S'))

        info_sheet = wb.create_sheet('导出信息')
        info_sheet.cell(row=1, column=1, value='筛选条件：').font = Font(bold=True)
        filter_row = 2
        for key, value in filters.items():
            info_sheet.cell(row=filter_row, column=1, value=key)
            info_sheet.cell(row=filter_row, column=2, value=str(value))
            filter_row += 1

        info_sheet.cell(row=filter_row + 1, column=1, value='生成时间：').font = Font(bold=True)
        info_sheet.cell(row=filter_row + 1, column=2,
                        value=timezone.now().strftime('%Y-%m-%d %H:%M:%S'))

        info_sheet.cell(row=filter_row + 2, column=1, value='操作人：').font = Font(bold=True)
        info_sheet.cell(row=filter_row + 2, column=2, value=request.user.username)

        file_name = f'续费跟进记录_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        DownloadRecord.objects.create(
            user=request.user,
            file_name=file_name,
            report_type=report_type,
            filters=filters,
            generated_by=request.user
        )

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'

        return response

    @action(detail=False, methods=['post'])
    def export_monthly_report(self, request):
        year = int(request.data.get('year', timezone.now().year))
        month = int(request.data.get('month', timezone.now().month))
        counselor_id = request.data.get('counselor')
        filters = {'year': year, 'month': month, 'counselor': counselor_id}

        from apps.followups.models import RenewalFollowUp
        from apps.courses.models import Enrollment

        first_day = datetime(year, month, 1)
        if month == 12:
            last_day = datetime(year + 1, 1, 1) - timedelta(seconds=1)
        else:
            last_day = datetime(year, month + 1, 1) - timedelta(seconds=1)

        follow_ups = RenewalFollowUp.objects.filter(
            created_at__gte=first_day,
            created_at__lte=last_day
        )

        if counselor_id:
            follow_ups = follow_ups.filter(counselor_id=counselor_id)

        total = follow_ups.count()
        completed = follow_ups.filter(status__in=['completed', 'renewed', 'closed']).count()
        renewed = follow_ups.filter(status='renewed').count()
        completion_rate = (completed / total * 100) if total > 0 else 0
        renewal_rate = (renewed / total * 100) if total > 0 else 0

        wb = openpyxl.Workbook()

        ws_summary = wb.active
        ws_summary.title = '月度复盘摘要'

        ws_summary.cell(row=1, column=1, value=f'{year}年{month}月 青少年培训续费跟进复盘报告').font = Font(
            bold=True, size=14)
        ws_summary.merge_cells('A1:D1')

        summary_data = [
            ['跟进总数', total],
            ['已完成数', completed],
            ['已续费数', renewed],
            ['完成率', f'{completion_rate:.2f}%'],
            ['续费率', f'{renewal_rate:.2f}%'],
        ]

        for i, (key, value) in enumerate(summary_data, 3):
            ws_summary.cell(row=i, column=1, value=key).font = Font(bold=True)
            ws_summary.cell(row=i, column=2, value=value)

        status_map = dict(RenewalFollowUp.STATUS_CHOICES)
        ws_status = wb.create_sheet('状态分布')
        ws_status.cell(row=1, column=1, value='状态').font = Font(bold=True)
        ws_status.cell(row=1, column=2, value='数量').font = Font(bold=True)
        status_counts = follow_ups.values('status').annotate(count=Count('id'))
        for i, item in enumerate(status_counts, 2):
            ws_status.cell(row=i, column=1, value=status_map.get(item['status'], item['status']))
            ws_status.cell(row=i, column=2, value=item['count'])

        ws_details = wb.create_sheet('跟进明细')
        headers = ['ID', '学员', '咨询师', '课程', '状态', '优先级', '原因', '创建时间']
        for col, header in enumerate(headers, 1):
            ws_details.cell(row=1, column=col, value=header).font = Font(bold=True)

        for row, item in enumerate(follow_ups, 2):
            ws_details.cell(row=row, column=1, value=item.id)
            ws_details.cell(row=row, column=2,
                            value=item.student.username if item.student else '')
            ws_details.cell(row=row, column=3,
                            value=item.counselor.username if item.counselor else '')
            ws_details.cell(row=row, column=4,
                            value=item.enrollment.course.name if item.enrollment and item.enrollment.course else '')
            ws_details.cell(row=row, column=5,
                            value=status_map.get(item.status, item.status))
            ws_details.cell(row=row, column=6, value=item.priority)
            ws_details.cell(row=row, column=7, value=item.reason)
            ws_details.cell(row=row, column=8,
                            value=item.created_at.strftime('%Y-%m-%d %H:%M:%S'))

        info_sheet = wb.create_sheet('导出信息')
        info_sheet.cell(row=1, column=1, value='筛选条件：').font = Font(bold=True)
        info_sheet.cell(row=2, column=1, value='年份')
        info_sheet.cell(row=2, column=2, value=year)
        info_sheet.cell(row=3, column=1, value='月份')
        info_sheet.cell(row=3, column=2, value=month)
        info_sheet.cell(row=4, column=1, value='咨询师')
        info_sheet.cell(row=4, column=2, value=counselor_id or '全部')

        info_sheet.cell(row=6, column=1, value='生成时间：').font = Font(bold=True)
        info_sheet.cell(row=6, column=2,
                        value=timezone.now().strftime('%Y-%m-%d %H:%M:%S'))

        info_sheet.cell(row=7, column=1, value='操作人：').font = Font(bold=True)
        info_sheet.cell(row=7, column=2, value=request.user.username)

        file_name = f'月度复盘报告_{year}年{month}月_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        DownloadRecord.objects.create(
            user=request.user,
            file_name=file_name,
            report_type='monthly_review',
            filters=filters,
            generated_by=request.user
        )

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{file_name}"'

        return response

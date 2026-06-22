import csv
import os
from datetime import datetime
from io import StringIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models import (
    AuditChecklist,
    SamplingRecord,
    RectificationPlan,
    Vendor,
    ExceptionOrder,
    SamplingStatus,
    EvidenceStatus,
    RiskLevel,
)
from app.schemas import SamplingCoverageResponse


class ExportService:
    SAMPLING_COVERAGE_DESCRIPTION = """
抽样覆盖率口径说明：
1. 覆盖率计算方式：已抽样检查清单数量 / 检查清单总数量 × 100%
2. 已抽样定义：检查清单下存在至少一条抽样记录（无论抽样记录状态）
3. 分类覆盖率：按检查清单的 category 字段分组统计
4. 统计时间：以导出时间为准
5. 注意事项：
   - 覆盖率仅反映抽样范围覆盖情况，不代表抽样深度
   - 同一检查清单多次抽样仅计算一次
   - 已删除的检查清单和抽样记录不计入统计
"""

    @staticmethod
    def _ensure_dirs():
        os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
        os.makedirs(settings.EXPORT_DIR, exist_ok=True)

    @staticmethod
    def _get_sampling_query_filters(filters: Optional[Dict[str, Any]] = None):
        conditions = []
        if filters:
            if filters.get("checklist_id"):
                conditions.append(SamplingRecord.checklist_id == filters["checklist_id"])
            if filters.get("status"):
                conditions.append(SamplingRecord.status == filters["status"])
            if filters.get("evidence_status"):
                conditions.append(SamplingRecord.evidence_status == filters["evidence_status"])
            if filters.get("start_date"):
                conditions.append(SamplingRecord.sampling_date >= filters["start_date"])
            if filters.get("end_date"):
                conditions.append(SamplingRecord.sampling_date <= filters["end_date"])
        return conditions

    @staticmethod
    def calculate_sampling_coverage(db: Session) -> SamplingCoverageResponse:
        total_checklists = db.query(AuditChecklist).count()

        sampled_checklist_ids = db.query(SamplingRecord.checklist_id).distinct().all()
        sampled_checklist_ids = [cid[0] for cid in sampled_checklist_ids]
        sampled_checklists = len(sampled_checklist_ids)

        coverage_rate = (sampled_checklists / total_checklists * 100) if total_checklists > 0 else 0.0

        by_category = {}
        categories = db.query(AuditChecklist.category).distinct().all()
        for (cat,) in categories:
            cat_total = db.query(AuditChecklist).filter(AuditChecklist.category == cat).count()
            cat_sampled_ids = db.query(SamplingRecord.checklist_id).join(
                AuditChecklist,
                SamplingRecord.checklist_id == AuditChecklist.id
            ).filter(AuditChecklist.category == cat).distinct().all()
            cat_sampled = len(cat_sampled_ids)
            cat_rate = (cat_sampled / cat_total * 100) if cat_total > 0 else 0.0
            by_category[cat] = {
                "total": cat_total,
                "sampled": cat_sampled,
                "coverage_rate": round(cat_rate, 2)
            }

        return SamplingCoverageResponse(
            total_checklists=total_checklists,
            sampled_checklists=sampled_checklists,
            coverage_rate=round(coverage_rate, 2),
            by_category=by_category,
            description=ExportService.SAMPLING_COVERAGE_DESCRIPTION.strip()
        )

    @staticmethod
    def export_sampling_records(
        db: Session,
        export_format: str = "excel",
        filters: Optional[Dict[str, Any]] = None
    ) -> Tuple[str, str]:
        ExportService._ensure_dirs()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        conditions = ExportService._get_sampling_query_filters(filters)
        query = db.query(SamplingRecord).join(
            AuditChecklist,
            SamplingRecord.checklist_id == AuditChecklist.id
        )
        if conditions:
            query = query.filter(*conditions)
        records = query.all()

        headers = [
            "抽样ID", "检查清单ID", "检查清单标题", "分类",
            "样本名称", "样本编码", "来源", "抽样日期",
            "抽样人", "状态", "证据状态", "创建时间"
        ]

        status_map = {
            "pending": "待审核",
            "reviewed": "已审核",
            "follow_up": "需跟进"
        }
        evidence_map = {
            "complete": "完整",
            "missing": "缺失",
            "partial": "部分"
        }

        rows = []
        for r in records:
            checklist = db.query(AuditChecklist).filter(AuditChecklist.id == r.checklist_id).first()
            rows.append([
                r.id,
                r.checklist_id,
                checklist.title if checklist else "",
                checklist.category if checklist else "",
                r.sample_name,
                r.sample_code,
                r.source or "",
                r.sampling_date.strftime("%Y-%m-%d") if r.sampling_date else "",
                r.sampled_by or "",
                status_map.get(r.status.value, r.status.value),
                evidence_map.get(r.evidence_status.value, r.evidence_status.value),
                r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else ""
            ])

        coverage = ExportService.calculate_sampling_coverage(db)

        if export_format == "csv":
            filename = f"sampling_records_{timestamp}.csv"
            filepath = os.path.join(settings.EXPORT_DIR, filename)
            with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                writer.writerows(rows)
                writer.writerow([])
                writer.writerow(["抽样覆盖率统计"])
                writer.writerow(["检查清单总数", coverage.total_checklists])
                writer.writerow(["已抽样检查清单数", coverage.sampled_checklists])
                writer.writerow(["总体覆盖率", f"{coverage.coverage_rate}%"])
                writer.writerow([])
                writer.writerow(["抽样覆盖口径说明"])
                for line in coverage.description.split("\n"):
                    writer.writerow([line])
            return filepath, filename

        else:
            filename = f"sampling_records_{timestamp}.xlsx"
            filepath = os.path.join(settings.EXPORT_DIR, filename)
            wb = Workbook()

            ws1 = wb.active
            ws1.title = "抽样记录"
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            for col, header in enumerate(headers, 1):
                cell = ws1.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row_idx, row in enumerate(rows, 2):
                for col_idx, val in enumerate(row, 1):
                    ws1.cell(row=row_idx, column=col_idx, value=val)

            for col in range(1, len(headers) + 1):
                ws1.column_dimensions[chr(64 + col) if col <= 26 else "A" + chr(64 + col - 26)].width = 18

            ws2 = wb.create_sheet("覆盖率统计")
            ws2["A1"] = "抽样覆盖率统计"
            ws2["A1"].font = Font(bold=True, size=14)
            ws2["A3"] = "检查清单总数"
            ws2["B3"] = coverage.total_checklists
            ws2["A4"] = "已抽样检查清单数"
            ws2["B4"] = coverage.sampled_checklists
            ws2["A5"] = "总体覆盖率"
            ws2["B5"] = f"{coverage.coverage_rate}%"

            ws2["A7"] = "按分类统计"
            ws2["A7"].font = Font(bold=True)
            ws2["A8"] = "分类"
            ws2["B8"] = "总数"
            ws2["C8"] = "已抽样"
            ws2["D8"] = "覆盖率"
            for cell in ["A8", "B8", "C8", "D8"]:
                ws2[cell].font = header_font
                ws2[cell].fill = header_fill

            row_idx = 9
            for cat, data in coverage.by_category.items():
                ws2.cell(row=row_idx, column=1, value=cat)
                ws2.cell(row=row_idx, column=2, value=data["total"])
                ws2.cell(row=row_idx, column=3, value=data["sampled"])
                ws2.cell(row=row_idx, column=4, value=f"{data['coverage_rate']}%")
                row_idx += 1

            ws3 = wb.create_sheet("口径说明")
            ws3["A1"] = "抽样覆盖口径说明"
            ws3["A1"].font = Font(bold=True, size=14)
            lines = coverage.description.split("\n")
            for i, line in enumerate(lines, 2):
                ws3.cell(row=i, column=1, value=line)

            wb.save(filepath)
            return filepath, filename

    @staticmethod
    def export_rectification_plans(
        db: Session,
        export_format: str = "excel",
        filters: Optional[Dict[str, Any]] = None
    ) -> Tuple[str, str]:
        ExportService._ensure_dirs()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        query = db.query(RectificationPlan)
        if filters:
            if filters.get("sampling_id"):
                query = query.filter(RectificationPlan.sampling_id == filters["sampling_id"])
            if filters.get("risk_level"):
                query = query.filter(RectificationPlan.risk_level == filters["risk_level"])
            if filters.get("status"):
                query = query.filter(RectificationPlan.status == filters["status"])
            if filters.get("vendor_id"):
                query = query.filter(RectificationPlan.vendor_id == filters["vendor_id"])
        plans = query.all()

        headers = [
            "计划ID", "抽样记录ID", "样本名称",
            "整改标题", "描述", "风险等级",
            "截止日期", "负责人", "供应商ID", "供应商名称",
            "状态", "创建时间", "更新时间"
        ]

        risk_map = {
            "low": "低",
            "medium": "中",
            "high": "高",
            "critical": "严重"
        }
        rect_status_map = {
            "not_started": "未启动",
            "in_progress": "进行中",
            "submitted": "已提交",
            "reviewed": "已审核",
            "closed": "已关闭"
        }

        rows = []
        for p in plans:
            sampling = db.query(SamplingRecord).filter(SamplingRecord.id == p.sampling_id).first()
            vendor = db.query(Vendor).filter(Vendor.id == p.vendor_id).first() if p.vendor_id else None
            rows.append([
                p.id,
                p.sampling_id,
                sampling.sample_name if sampling else "",
                p.title,
                p.description or "",
                risk_map.get(p.risk_level.value, p.risk_level.value),
                p.deadline.strftime("%Y-%m-%d") if p.deadline else "",
                p.responsible_person or "",
                p.vendor_id or "",
                vendor.name if vendor else "",
                rect_status_map.get(p.status.value, p.status.value),
                p.created_at.strftime("%Y-%m-%d %H:%M:%S") if p.created_at else "",
                p.updated_at.strftime("%Y-%m-%d %H:%M:%S") if p.updated_at else ""
            ])

        if export_format == "csv":
            filename = f"rectification_plans_{timestamp}.csv"
            filepath = os.path.join(settings.EXPORT_DIR, filename)
            with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                writer.writerows(rows)
            return filepath, filename

        else:
            filename = f"rectification_plans_{timestamp}.xlsx"
            filepath = os.path.join(settings.EXPORT_DIR, filename)
            wb = Workbook()
            ws = wb.active
            ws.title = "整改计划"
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
            for row_idx, row in enumerate(rows, 2):
                for col_idx, val in enumerate(row, 1):
                    ws.cell(row=row_idx, column=col_idx, value=val)
            for col in range(1, len(headers) + 1):
                ws.column_dimensions[chr(64 + col) if col <= 26 else "A" + chr(64 + col - 26)].width = 18
            wb.save(filepath)
            return filepath, filename

    @staticmethod
    def export_exception_orders(
        db: Session,
        export_format: str = "excel",
        filters: Optional[Dict[str, Any]] = None
    ) -> Tuple[str, str]:
        ExportService._ensure_dirs()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        query = db.query(ExceptionOrder)
        if filters:
            if filters.get("sampling_id"):
                query = query.filter(ExceptionOrder.sampling_id == filters["sampling_id"])
            if filters.get("exception_type"):
                query = query.filter(ExceptionOrder.exception_type == filters["exception_type"])
            if filters.get("status"):
                query = query.filter(ExceptionOrder.status == filters["status"])
        orders = query.all()

        headers = [
            "异常单ID", "抽样记录ID", "样本名称",
            "异常类型", "影响范围", "负责人",
            "根本原因", "处理结果", "状态",
            "创建时间", "更新时间"
        ]

        type_map = {
            "evidence_missing": "证据缺失",
            "non_compliance": "不合规",
            "other": "其他"
        }
        exc_status_map = {
            "open": "待处理",
            "processing": "处理中",
            "closed": "已关闭"
        }

        rows = []
        for e in orders:
            sampling = db.query(SamplingRecord).filter(SamplingRecord.id == e.sampling_id).first()
            rows.append([
                e.id,
                e.sampling_id,
                sampling.sample_name if sampling else "",
                type_map.get(e.exception_type.value, e.exception_type.value),
                e.impact_scope or "",
                e.responsible_person or "",
                e.root_cause or "",
                e.handling_result or "",
                exc_status_map.get(e.status.value, e.status.value),
                e.created_at.strftime("%Y-%m-%d %H:%M:%S") if e.created_at else "",
                e.updated_at.strftime("%Y-%m-%d %H:%M:%S") if e.updated_at else ""
            ])

        if export_format == "csv":
            filename = f"exception_orders_{timestamp}.csv"
            filepath = os.path.join(settings.EXPORT_DIR, filename)
            with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                writer.writerows(rows)
            return filepath, filename

        else:
            filename = f"exception_orders_{timestamp}.xlsx"
            filepath = os.path.join(settings.EXPORT_DIR, filename)
            wb = Workbook()
            ws = wb.active
            ws.title = "异常单"
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="FF6347", end_color="FF6347", fill_type="solid")
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
            for row_idx, row in enumerate(rows, 2):
                for col_idx, val in enumerate(row, 1):
                    ws.cell(row=row_idx, column=col_idx, value=val)
            for col in range(1, len(headers) + 1):
                ws.column_dimensions[chr(64 + col) if col <= 26 else "A" + chr(64 + col - 26)].width = 20
            wb.save(filepath)
            return filepath, filename

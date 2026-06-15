import io
import os
import json
from datetime import datetime, date
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException, Query, BackgroundTasks
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, and_, extract
from sqlalchemy.orm import selectinload
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from app.database import get_db
from app.models import (
    ReviewApplication, ReviewStatus, User, UserRole, ReportDownload,
    Classroom, ClassroomSchedule, AuditLog, AuditAction,
    Student, Score, Course,
)
from app.schemas import ReportDownloadResponse, ClassroomUtilization, MonthlyReviewSummary
from app.security import get_current_user, require_roles

router = APIRouter()

REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "generated_reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


def _enrich_report(r: ReportDownload) -> dict:
    r_dict = r.__dict__
    if r.generated_by:
        r_dict["generated_by_name"] = r.generated_by.full_name
    return r_dict


@router.get("/classroom-utilization", response_model=list[ClassroomUtilization])
async def get_classroom_utilization(
    year: int = Query(..., description="年份"),
    month: int = Query(..., description="月份 1-12"),
    building: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    query = select(Classroom)
    if building:
        query = query.where(Classroom.building == building)
    classrooms = (await db.execute(query)).scalars().all()

    result = []
    TOTAL_PERIODS_PER_DAY = 10
    import calendar
    days_in_month = calendar.monthrange(year, month)[1]

    for cr in classrooms:
        schedules = (await db.execute(
            select(ClassroomSchedule).where(
                ClassroomSchedule.classroom_id == cr.id,
                extract("year", ClassroomSchedule.date) == year,
                extract("month", ClassroomSchedule.date) == month,
            )
        )).scalars().all()

        used_periods = sum((s.period_end - s.period_start + 1) for s in schedules)
        total_periods = TOTAL_PERIODS_PER_DAY * days_in_month
        utilization_rate = round(used_periods / total_periods * 100, 2) if total_periods > 0 else 0

        attendances = [s.actual_attendance for s in schedules if s.actual_attendance]
        avg_attendance = round(sum(attendances) / len(attendances), 2) if attendances else None

        result.append(ClassroomUtilization(
            classroom_id=cr.id,
            building=cr.building,
            room_no=cr.room_no,
            capacity=cr.capacity,
            total_periods=total_periods,
            used_periods=used_periods,
            utilization_rate=utilization_rate,
            avg_attendance=avg_attendance,
        ))
    return result


@router.get("/monthly-summary", response_model=list[MonthlyReviewSummary])
async def get_monthly_summary(
    year: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    query = (
        select(
            func.extract("year", ReviewApplication.applied_at).label("y"),
            func.extract("month", ReviewApplication.applied_at).label("m"),
            func.count(ReviewApplication.id).label("total"),
        )
        .group_by("y", "m")
        .order_by("y", "m")
    )
    if year:
        query = query.where(extract("year", ReviewApplication.applied_at) == year)
    months_data = (await db.execute(query)).all()

    result = []
    statuses = [ReviewStatus.PENDING, ReviewStatus.UNDER_REVIEW, ReviewStatus.MATERIALS_MISSING,
                ReviewStatus.APPROVED, ReviewStatus.REJECTED, ReviewStatus.CLOSED]

    for y, m, total in months_data:
        month_str = f"{int(y):04d}-{int(m):02d}"
        counts = {}
        for s in statuses:
            c = (await db.execute(
                select(func.count(ReviewApplication.id)).where(
                    extract("year", ReviewApplication.applied_at) == y,
                    extract("month", ReviewApplication.applied_at) == m,
                    ReviewApplication.status == s,
                )
            )).scalar_one()
            counts[s.value] = c

        days_data = (await db.execute(
            select(ReviewApplication.applied_at, ReviewApplication.closed_at).where(
                extract("year", ReviewApplication.applied_at) == y,
                extract("month", ReviewApplication.applied_at) == m,
                ReviewApplication.closed_at.isnot(None),
            )
        )).all()
        avg_days = None
        if days_data:
            days_list = [(c - o).days for o, c in days_data if c and o]
            if days_list:
                avg_days = round(sum(days_list) / len(days_list), 2)

        result.append(MonthlyReviewSummary(
            year_month=month_str,
            total_reviews=total,
            pending_count=counts["pending"],
            under_review_count=counts["under_review"],
            materials_missing_count=counts["materials_missing"],
            approved_count=counts["approved"],
            rejected_count=counts["rejected"],
            closed_count=counts["closed"],
            avg_processing_days=avg_days,
        ))
    return result


def _generate_excel_report(report_type: str, data: list, filter_criteria: dict, generated_by: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = report_type

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    meta_font = Font(bold=True, size=11)
    meta_val_font = Font(size=11)

    if not data:
        ws.append(["无数据"])
    else:
        headers = list(data[0].keys() if isinstance(data[0], dict) else data[0].model_dump().keys())
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row_data in data:
            row_dict = row_data if isinstance(row_data, dict) else row_data.model_dump()
            ws.append([row_dict.get(h, "") for h in headers])

    ws.append([])
    ws.append(["报表元信息"])
    ws.append(["生成者", generated_by or "未知"])
    ws.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws.append([])
    ws.append(["筛选口径"])
    filter_items = list(filter_criteria.items())
    for k, v in filter_items:
        ws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else str(v)])

    for row in ws.iter_rows(min_row=ws.max_row - len(filter_items) - 4, max_row=ws.max_row, min_col=1, max_col=1):
        for cell in row:
            if cell.value and not isinstance(cell.value, (int, float)):
                cell.font = meta_font

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@router.post("/generate/{report_type}")
async def generate_report(
    report_type: str,
    filter_criteria: dict,
    current_user: User = Depends(require_roles(UserRole.ADMIN, UserRole.STUDENT_AFFAIRS, UserRole.AUDITOR)),
    db: AsyncSession = Depends(get_db),
):
    valid_types = ["classroom_utilization", "monthly_summary", "review_details", "advisor_quota"]
    if report_type not in valid_types:
        raise HTTPException(status_code=400, detail=f"无效的报表类型，支持: {', '.join(valid_types)}")

    report = ReportDownload(
        report_name=f"{report_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}",
        report_type=report_type,
        filter_criteria=filter_criteria,
        generated_by_id=current_user.id,
        status="completed",
    )
    db.add(report)
    await db.flush()

    data = []
    if report_type == "classroom_utilization":
        data_dicts = await get_classroom_utilization(
            filter_criteria.get("year", datetime.now().year),
            filter_criteria.get("month", datetime.now().month),
            filter_criteria.get("building"),
            current_user, db,
        )
        data = [d.model_dump() for d in data_dicts]
    elif report_type == "monthly_summary":
        data_dicts = await get_monthly_summary(filter_criteria.get("year"), current_user, db)
        data = [d.model_dump() for d in data_dicts]
    elif report_type == "review_details":
        reviews = (await db.execute(
            select(ReviewApplication).options(selectinload(ReviewApplication.student), selectinload(ReviewApplication.course))
        )).scalars().all()
        data = []
        for r in reviews:
            data.append({
                "申请编号": r.application_no,
                "学生": r.student.name if r.student else "",
                "课程": r.course.course_name if r.course else "",
                "当前成绩": r.current_score,
                "状态": r.status.value,
                "申请时间": r.applied_at.strftime("%Y-%m-%d %H:%M") if r.applied_at else "",
            })
    elif report_type == "advisor_quota":
        from app.models import AdvisorQuota
        quotas = (await db.execute(
            select(AdvisorQuota).options(selectinload(AdvisorQuota.advisor))
        )).scalars().all()
        data = []
        for q in quotas:
            data.append({
                "导师": q.advisor.full_name if q.advisor else "",
                "学期": q.semester,
                "最大名额": q.max_quota,
                "已分配": q.current_assigned,
                "剩余名额": q.max_quota - q.current_assigned,
            })

    file_bytes = _generate_excel_report(report_type, data, filter_criteria, generated_by=current_user.full_name)
    file_name = f"{report.report_name}.xlsx"
    file_path = os.path.join(REPORTS_DIR, file_name)
    with open(file_path, "wb") as f:
        f.write(file_bytes)
    report.file_path = file_name
    await db.commit()
    await db.refresh(report)

    return {
        "report_id": report.id,
        "file_name": file_name,
        "generated_by": current_user.full_name,
        "filter_criteria": filter_criteria,
        "message": "报表生成成功",
    }


@router.get("/download/{report_id}")
async def download_report(
    report_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(ReportDownload).options(selectinload(ReportDownload.generated_by)).where(ReportDownload.id == report_id))
    report = result.scalar_one_or_none()
    if not report:
        raise HTTPException(status_code=404, detail="报表不存在")
    if not report.file_path:
        raise HTTPException(status_code=400, detail="报表文件不存在")

    report.download_count += 1
    log = AuditLog(
        action=AuditAction.DOWNLOAD,
        entity_type="report",
        entity_id=report.id,
        operator_id=current_user.id,
        new_values={
            "report_name": report.report_name,
            "report_type": report.report_type,
            "filter_criteria": report.filter_criteria,
        },
    )
    db.add(log)
    await db.commit()

    full_path = os.path.join(REPORTS_DIR, report.file_path)
    if not os.path.exists(full_path):
        raise HTTPException(status_code=404, detail="文件已过期")

    def iter_file():
        with open(full_path, "rb") as f:
            while chunk := f.read(8192):
                yield chunk

    headers = {
        "Content-Disposition": f'attachment; filename="{report.file_path}"',
        "X-Report-Name": quote(report.report_name, safe=""),
        "X-Generated-By": quote(report.generated_by.full_name, safe="") if report.generated_by else "",
        "X-Filter-Criteria": quote(json.dumps(report.filter_criteria, ensure_ascii=False), safe=""),
    }
    return StreamingResponse(
        iter_file(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.get("", response_model=dict)
async def list_reports(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    report_type: Optional[str] = None,
    generated_by_id: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    query = select(ReportDownload).options(selectinload(ReportDownload.generated_by))
    count_query = select(func.count(ReportDownload.id))

    if report_type:
        query = query.where(ReportDownload.report_type == report_type)
        count_query = count_query.where(ReportDownload.report_type == report_type)
    if generated_by_id:
        query = query.where(ReportDownload.generated_by_id == generated_by_id)
        count_query = count_query.where(ReportDownload.generated_by_id == generated_by_id)

    total = (await db.execute(count_query)).scalar_one()
    query = query.order_by(ReportDownload.generated_at.desc()).offset((page - 1) * page_size).limit(page_size)
    result = await db.execute(query)
    reports = result.scalars().all()

    return {
        "data": [_enrich_report(r) for r in reports],
        "pagination": {
            "page": page,
            "page_size": page_size,
            "total": total,
            "total_pages": (total + page_size - 1) // page_size,
        },
    }


@router.get("/{report_id}/info", response_model=ReportDownloadResponse)
async def get_report_info(
    report_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(ReportDownload).options(selectinload(ReportDownload.generated_by)).where(ReportDownload.id == report_id))
    report = result.scalar_one_or_none()
    if not report:
        raise HTTPException(status_code=404, detail="报表不存在")
    return _enrich_report(report)

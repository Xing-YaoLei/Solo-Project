from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd

from config import Config
from app.data.queries import (
    get_project_funnel,
    get_reconciliation_trend,
    get_contract_attachments,
    get_attachment_type_stats,
    get_document_details,
    get_approval_abnormal,
    get_payment_cycle_data,
    get_project_list
)
from app.data.services import build_reconciliation_trend_df, aggregate_monthly_diff


PAYMENT_CYCLE_DEFINITION = """
【回款周期口径说明】

1. 回款周期定义：
   - 指从合同签订日期到实际回款到账日期之间的自然日天数
   - 公式：回款周期(天) = 实际回款日期 - 合同签订日期
   - 若实际回款日期为空，则回款周期显示为"-"（未完成）

2. 逾期天数定义：
   - 指实际回款日期超出计划回款日期的自然日天数
   - 公式：逾期天数(天) = 实际回款日期 - 计划回款日期
   - 提前回款的，逾期天数显示为负数
   - 若计划日期或实际日期为空，则逾期天数显示为0

3. 回款阶段划分：
   - 定金：合同签订时支付（一般为合同金额的10%~30%）
   - 开工款：开工前支付（一般为合同金额的30%~40%）
   - 进度款：按施工进度节点支付（一般分2~3期，每期20%~30%）
   - 竣工款：竣工验收合格后支付（一般为合同金额的5%~10%）
   - 质保金：质保期满后支付（一般为合同金额的3%~5%，质保期通常1~2年）

4. 数据统计范围：
   - 包含所有已录入系统的合同及回款记录
   - 不包含已作废/已取消的合同
   - 金额单位为人民币元（¥）

5. 异常情况说明：
   - 回款状态为"未收款"的记录不参与回款周期均值计算
   - 同一合同的多笔回款按阶段分别计算周期
"""


def export_full_report() -> str:
    Config.ensure_dirs()
    timestamp = date.today().strftime("%Y%m%d")
    filename = f"家装量房报价漏斗报表_{timestamp}.xlsx"
    filepath = Config.EXPORT_DIR / filename

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_cover_sheet(wb)
    _write_definition_sheet(wb)
    _write_sheet_from_df(wb, "漏斗概览", get_project_funnel())

    rec_df = build_reconciliation_trend_df(months=12)
    _write_sheet_from_df(wb, "对账差异明细", rec_df)
    if not rec_df.empty:
        monthly = aggregate_monthly_diff(rec_df)
        _write_sheet_from_df(wb, "对账差异月度汇总", monthly)

    _write_sheet_from_df(wb, "合同附件构成", get_attachment_type_stats())
    _write_sheet_from_df(wb, "合同附件明细", get_contract_attachments())
    _write_sheet_from_df(wb, "报价单明细", get_document_details("quotation"))
    _write_sheet_from_df(wb, "采购单明细", get_document_details("purchase"))
    _write_sheet_from_df(wb, "审批异常", get_approval_abnormal())
    _write_sheet_from_df(wb, "回款周期", get_payment_cycle_data())
    _write_sheet_from_df(wb, "项目清单", get_project_list())

    wb.save(filepath)
    return str(filepath)


def export_single_sheet(sheet_type: str, df: pd.DataFrame) -> str:
    Config.ensure_dirs()
    timestamp = date.today().strftime("%Y%m%d_%H%M%S")
    filename = f"{sheet_type}_{timestamp}.xlsx"
    filepath = Config.EXPORT_DIR / filename

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_type
    _write_definition_header(ws, sheet_type)
    start_row = ws.max_row + 2
    _write_df_to_worksheet(ws, df, start_row)
    wb.save(filepath)
    return str(filepath)


def _write_cover_sheet(wb: openpyxl.Workbook):
    ws = wb.create_sheet("封面说明", 0)

    title_font = Font(name="微软雅黑", size=20, bold=True, color="1F4E79")
    header_font = Font(name="微软雅黑", size=12, bold=True, color="2E75B6")
    normal_font = Font(name="微软雅黑", size=10)

    ws["A1"] = "家装工地量房报价漏斗报表"
    ws["A1"].font = title_font
    ws.merge_cells("A1:F1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A3"] = f"报表生成日期：{date.today().strftime('%Y年%m月%d日')}"
    ws["A3"].font = header_font
    ws["A4"] = f"报表范围：全量项目数据（含量房、报价、合同、采购、回款全流程）"
    ws["A4"].font = normal_font
    ws["A5"] = f"数据口径：详见「回款周期口径」工作表"
    ws["A5"].font = normal_font

    ws["A7"] = "报表结构说明"
    ws["A7"].font = header_font
    ws.merge_cells("A7:F7")

    sheets_info = [
        ("漏斗概览", "量房→报价→签单→施工→竣工各阶段项目转化漏斗"),
        ("对账差异明细", "采购单预算金额与实际金额的逐笔差异明细"),
        ("对账差异月度汇总", "按月份汇总的对账差异金额与差异率"),
        ("合同附件构成", "合同附件类型、数量、大小统计"),
        ("合同附件明细", "每份合同的附件明细列表"),
        ("报价单明细", "所有报价单的分项明细（溯源基础）"),
        ("采购单明细", "所有采购单的材料明细（溯源基础）"),
        ("审批异常", "超时效、被驳回等审批异常记录"),
        ("回款周期", "各项目各回款阶段的计划/实际周期明细"),
        ("项目清单", "所有项目基础信息索引"),
        ("回款周期口径", "回款周期计算口径、异常情况说明"),
    ]
    for i, (name, desc) in enumerate(sheets_info):
        row = 8 + i
        ws.cell(row=row, column=1, value=name).font = Font(name="微软雅黑", size=10, bold=True)
        ws.cell(row=row, column=2, value=desc).font = normal_font
        ws.merge_cells(f"B{row}:F{row}")

    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 22


def _write_definition_sheet(wb: openpyxl.Workbook):
    ws = wb.create_sheet("回款周期口径")
    header_font = Font(name="微软雅黑", size=14, bold=True, color="C00000")
    title_font = Font(name="微软雅黑", size=11, bold=True, color="1F4E79")
    normal_font = Font(name="微软雅黑", size=10)

    ws["A1"] = "回款周期与对账数据口径说明"
    ws["A1"].font = header_font
    ws.merge_cells("A1:H1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    row = 3
    for line in PAYMENT_CYCLE_DEFINITION.strip().split("\n"):
        line = line.strip()
        if not line:
            row += 1
            continue
        if line.startswith("【") and line.endswith("】"):
            ws.cell(row=row, column=1, value=line).font = title_font
            ws.merge_cells(f"A{row}:H{row}")
            ws.cell(row=row, column=1).fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        elif line[:1].isdigit() and ". " in line[:4]:
            ws.cell(row=row, column=1, value=line).font = Font(name="微软雅黑", size=10, bold=True)
            ws.merge_cells(f"A{row}:H{row}")
        else:
            ws.cell(row=row, column=2, value=line).font = normal_font
            ws.merge_cells(f"B{row}:H{row}")
        row += 1

    ws.column_dimensions["A"].width = 8
    for col in range(2, 9):
        ws.column_dimensions[get_column_letter(col)].width = 20


def _write_sheet_from_df(wb: openpyxl.Workbook, title: str, df: pd.DataFrame):
    ws = wb.create_sheet(title)
    _write_definition_header(ws, title)
    start_row = ws.max_row + 2
    _write_df_to_worksheet(ws, df, start_row)


def _write_definition_header(ws, sheet_name: str):
    header_map = {
        "对账差异明细": [
            "对账差异 = 预算金额 - 实际金额；差异率 = 差异金额 / 预算金额 × 100%",
            "正差异表示有结余（预算未花完），负差异表示超预算"
        ],
        "对账差异月度汇总": [
            "按月度汇总采购差异；月差异率 = 月度差异总额 / 月度预算总额 × 100%"
        ],
        "审批异常": [
            "判定规则：1) 审批通过时间超期望时效；2) 审批被驳回；3) 提交超过7天未处理"
        ],
        "回款周期": [
            "回款周期 = 实际回款日期 - 合同签订日期；逾期天数 = 实际回款日期 - 计划回款日期",
            "详细口径说明请参阅「回款周期口径」工作表"
        ],
    }

    title_font = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
    note_font = Font(name="微软雅黑", size=9, color="595959", italic=True)

    ws["A1"] = f"【{sheet_name}】数据明细"
    ws["A1"].font = title_font

    notes = header_map.get(sheet_name, [])
    for i, note in enumerate(notes):
        ws.cell(row=2 + i, column=1, value=f"※ {note}").font = note_font


def _write_df_to_worksheet(ws, df: pd.DataFrame, start_row: int = 1):
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_font = Font(name="微软雅黑", size=9)
    cell_align = Alignment(vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    if df is None or df.empty:
        ws.cell(row=start_row, column=1, value="(暂无数据)").font = Font(italic=True, color="808080")
        return

    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for r_idx, row in enumerate(df.itertuples(index=False), start_row + 1):
        for c_idx, value in enumerate(row, 1):
            val = value
            if pd.isna(val) if hasattr(pd, "isna") else val is None:
                val = ""
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = cell_font
            cell.alignment = cell_align
            cell.border = thin_border
            if c_idx >= 5 and isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'

    for col_idx in range(1, len(df.columns) + 1):
        max_len = 12
        for row in range(start_row, start_row + min(len(df) + 1, 200)):
            v = ws.cell(row=row, column=col_idx).value
            if v:
                max_len = max(max_len, min(len(str(v)) + 2, 40))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len

    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)

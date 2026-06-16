from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


class ExportManager:
    EXPORT_DEFINITIONS = {
        "elders": {
            "title": "老人档案清单",
            "caliber": "数据口径说明：\n1. 统计范围：在院所有老人档案\n2. 健康状态：stable-稳定, monitoring-观察中, critical-危重\n3. 护理等级：basic-基础护理, intermediate-中级护理, advanced-高级护理, special-特护\n4. 数据更新时间：{export_time}",
            "columns": [
                ("姓名", "name"),
                ("性别", "gender"),
                ("出生日期", "birth_date"),
                ("身份证号", "id_card"),
                ("联系电话", "phone"),
                ("紧急联系人", "emergency_contact"),
                ("紧急联系电话", "emergency_phone"),
                ("健康状态", "health_status"),
                ("护理等级", "care_level"),
                ("房间号", "room_number"),
                ("床号", "bed_number"),
                ("入院日期", "admission_date"),
                ("状态", "status"),
            ]
        },
        "medications": {
            "title": "用药清单",
            "caliber": "数据口径说明：\n1. 统计范围：所有在用及历史用药记录\n2. 用药状态：active-在用, discontinued-已停用, completed-已完成\n3. 频次说明：qd-每日一次, bid-每日两次, tid-每日三次, qid-每日四次\n4. 数据更新时间：{export_time}",
            "columns": [
                ("老人姓名", "elder_name"),
                ("药品名称", "drug_name"),
                ("通用名", "generic_name"),
                ("剂量", "dosage"),
                ("频次", "frequency"),
                ("给药途径", "route"),
                ("开始日期", "start_date"),
                ("结束日期", "end_date"),
                ("开方医生", "prescribing_doctor"),
                ("用药目的", "purpose"),
                ("状态", "status"),
            ]
        },
        "visit_records": {
            "title": "探访记录",
            "caliber": "数据口径说明：\n1. 统计范围：所有探访记录\n2. 探访类型：routine-例行探访, family-家属探访, medical-医疗探访, other-其他\n3. 时长单位：分钟\n4. 数据更新时间：{export_time}",
            "columns": [
                ("老人姓名", "elder_name"),
                ("探访日期", "visit_date"),
                ("探访时间", "visit_time"),
                ("探访时长(分钟)", "visit_duration"),
                ("探访类型", "visit_type"),
                ("探访人姓名", "visitor_name"),
                ("与老人关系", "visitor_relation"),
                ("身体状况", "physical_condition"),
                ("精神状况", "mental_condition"),
                ("老人情绪", "elder_mood"),
                ("状态", "status"),
            ]
        },
        "activities": {
            "title": "活动签到记录",
            "caliber": "数据口径说明：\n1. 统计范围：所有康复活动及签到记录\n2. 活动类型：rehabilitation-康复训练, entertainment-文娱活动, exercise-健身锻炼, education-健康教育\n3. 参与状态：signed_in-已签到, participated-已参与, absent-缺席, left_early-早退\n4. 数据更新时间：{export_time}",
            "columns": [
                ("活动名称", "activity_name"),
                ("活动类型", "activity_type"),
                ("活动日期", "activity_date"),
                ("开始时间", "start_time"),
                ("结束时间", "end_time"),
                ("活动地点", "location"),
                ("老人姓名", "elder_name"),
                ("签到时间", "sign_in_time"),
                ("签退时间", "sign_out_time"),
                ("参与状态", "participation_status"),
                ("表现评分", "performance_rating"),
            ]
        },
        "risk_events": {
            "title": "风险事件记录",
            "caliber": "数据口径说明：\n1. 统计范围：所有上报的风险事件\n2. 事件类型：fall-跌倒, pressure_ulcer-压疮, medication_error-用药错误,走失-wandering, other-其他\n3. 事件等级：general-一般, serious-严重, critical-危重\n4. 状态：reported-已上报, investigating-调查中, handled-已处理, closed-已结案\n5. 数据更新时间：{export_time}",
            "columns": [
                ("老人姓名", "elder_name"),
                ("事件类型", "event_type"),
                ("事件等级", "event_level"),
                ("事件日期", "event_date"),
                ("事件时间", "event_time"),
                ("发生地点", "location"),
                ("事件描述", "description"),
                ("伤情", "injuries"),
                ("即时措施", "immediate_measures"),
                ("状态", "status"),
                ("处理结果", "handling_result"),
            ]
        },
        "incident_orders": {
            "title": "异常单记录",
            "caliber": "数据口径说明：\n1. 统计范围：所有严重异常事件单（含跌倒、走失等）\n2. 严重程度：minor-轻微, moderate-中度, serious-严重, critical-危重\n3. 状态：pending-待处理, processing-处理中, closed-已结案\n4. 包含字段：影响范围、责任归属、处理结果\n5. 数据更新时间：{export_time}",
            "columns": [
                ("异常单号", "order_no"),
                ("老人姓名", "elder_name"),
                ("异常类型", "incident_type"),
                ("严重程度", "severity"),
                ("发生日期", "incident_date"),
                ("发生时间", "incident_time"),
                ("发生地点", "location"),
                ("影响范围", "impact_scope"),
                ("责任归属", "responsibility"),
                ("责任人", "responsible_person"),
                ("处理结果", "handling_result"),
                ("预防措施", "preventive_measures"),
                ("状态", "status"),
            ]
        }
    }

    @classmethod
    def export_to_excel(cls, export_type: str, data: list, extra_context: dict = None) -> BytesIO:
        if export_type not in cls.EXPORT_DEFINITIONS:
            raise ValueError(f"不支持的导出类型: {export_type}")

        definition = cls.EXPORT_DEFINITIONS[export_type]
        export_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        wb = Workbook()
        ws = wb.active
        ws.title = definition["title"]

        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        title_font = Font(bold=True, size=14)
        caliber_font = Font(size=10, color="666666")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(definition["columns"]))
        ws.cell(row=1, column=1, value=definition["title"]).font = title_font
        ws.cell(row=1, column=1).alignment = center_align

        caliber_text = definition["caliber"].format(export_time=export_time)
        caliber_rows = len(caliber_text.split('\n'))
        ws.merge_cells(start_row=2, start_column=1, end_row=1 + caliber_rows,
                       end_column=len(definition["columns"]))
        ws.cell(row=2, column=1, value=caliber_text).font = caliber_font
        ws.cell(row=2, column=1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        header_row = 2 + caliber_rows
        for col_idx, (col_name, _) in enumerate(definition["columns"], 1):
            cell = ws.cell(row=header_row, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for row_idx, item in enumerate(data, header_row + 1):
            for col_idx, (_, field) in enumerate(definition["columns"], 1):
                value = item.get(field, "") if isinstance(item, dict) else getattr(item, field, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = left_align if col_idx > 2 else center_align
                cell.border = thin_border

        for col_idx in range(1, len(definition["columns"]) + 1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 15

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def get_export_caliber(cls, export_type: str) -> str:
        if export_type not in cls.EXPORT_DEFINITIONS:
            raise ValueError(f"不支持的导出类型: {export_type}")
        export_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        return cls.EXPORT_DEFINITIONS[export_type]["caliber"].format(export_time=export_time)

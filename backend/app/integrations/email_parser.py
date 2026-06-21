import hashlib
import imaplib
import logging
import os
import random
import re
from datetime import datetime, timedelta
from enum import StrEnum
from typing import Any
from uuid import UUID, uuid4

try:
    import mailparser
    MAILPARSER_AVAILABLE = True
except ImportError:
    MAILPARSER_AVAILABLE = False
    mailparser = None

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    openpyxl = None

from pydantic import BaseModel, ConfigDict, Field
from pydantic_settings import SettingsConfigDict
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config.settings import settings
from app.integrations.base_client import (
    BaseIntegrationClient,
    IntegrationConfig,
    IntegrationConnectionError,
    IntegrationError,
)
from app.models.email_attachment import EmailAttachment
from app.models.invoice import Invoice, InvoiceSource, InvoiceStatus
from app.schemas.case import InvoiceCreate, InvoiceItemCreate

logger = logging.getLogger(__name__)


class AttachmentType(StrEnum):
    INVOICE = "invoice"
    BILL = "bill"
    CONTRACT = "contract"
    RECEIPT = "receipt"
    STATEMENT = "statement"
    OTHER = "other"


class EmailConfig(IntegrationConfig):
    model_config = SettingsConfigDict(env_prefix="EMAIL_", extra="ignore")

    imap_host: str | None = Field(default=None)
    imap_port: int = Field(default=993)
    imap_username: str | None = Field(default=None)
    imap_password: str | None = Field(default=None)
    imap_folder: str = Field(default="INBOX")
    use_ssl: bool = Field(default=True)
    attachment_dir: str = Field(default="./attachments")


class ParsedAttachment(BaseModel):
    model_config = ConfigDict(extra="allow")

    file_name: str
    file_path: str
    file_size: int
    content_type: str
    hash: str
    attachment_type: AttachmentType
    content: str | None = None
    case_no: str | None = None
    invoice_no: str | None = None
    amount: float | None = None


class ParsedEmail(BaseModel):
    model_config = ConfigDict(extra="allow")

    message_id: str
    subject: str | None
    sender: str | None
    received_at: datetime
    body: str | None
    attachments: list[ParsedAttachment]


class EmailParser(BaseIntegrationClient[ParsedEmail]):
    def __init__(
        self,
        config: EmailConfig | None = None,
        db: AsyncSession | None = None,
    ) -> None:
        super().__init__(
            config=config or EmailConfig(),
            db=db,
            source_system="email",
        )
        self._imap_connection: imaplib.IMAP4_SSL | imaplib.IMAP4 | None = None
        self._processed_hashes: set[str] = set()

    async def _connect_imap(self) -> None:
        if self._imap_connection:
            return

        config = self.config
        if not isinstance(config, EmailConfig):
            config = EmailConfig(**config.model_dump())

        if not config.imap_host or not config.imap_username or not config.imap_password:
            raise IntegrationError("IMAP credentials not configured")

        try:
            if config.use_ssl:
                self._imap_connection = imaplib.IMAP4_SSL(
                    config.imap_host, config.imap_port, timeout=config.timeout
                )
            else:
                self._imap_connection = imaplib.IMAP4(
                    config.imap_host, config.imap_port, timeout=config.timeout
                )

            self._imap_connection.login(config.imap_username, config.imap_password)
            self._imap_connection.select(config.imap_folder)
            logger.info(f"Connected to IMAP server: {config.imap_host}")
        except Exception as e:
            raise IntegrationConnectionError(f"Failed to connect to IMAP: {e}") from e

    async def _disconnect_imap(self) -> None:
        if self._imap_connection:
            try:
                self._imap_connection.close()
                self._imap_connection.logout()
            except Exception as e:
                logger.warning(f"Error disconnecting from IMAP: {e}")
            finally:
                self._imap_connection = None

    async def fetch_data(
        self,
        since: datetime | None = None,
        limit: int = 50,
        unread_only: bool = True,
    ) -> list[ParsedEmail]:
        if self.use_mock:
            return await self._generate_mock_data(limit=limit)

        async def _fetch() -> list[ParsedEmail]:
            await self._connect_imap()
            return await self._fetch_emails(since=since, limit=limit, unread_only=unread_only)

        try:
            return await self._retry(_fetch)
        except Exception as e:
            logger.warning(f"Failed to fetch emails, falling back to mock: {e}")
            return await self._generate_mock_data(limit=limit)
        finally:
            await self._disconnect_imap()

    async def _fetch_emails(
        self,
        since: datetime | None = None,
        limit: int = 50,
        unread_only: bool = True,
    ) -> list[ParsedEmail]:
        if not self._imap_connection:
            raise IntegrationError("Not connected to IMAP server")

        search_criteria = []
        if unread_only:
            search_criteria.append("UNSEEN")
        if since:
            since_str = since.strftime("%d-%b-%Y")
            search_criteria.append(f"SINCE {since_str}")

        search_query = " ".join(search_criteria) if search_criteria else "ALL"

        try:
            status, messages = self._imap_connection.search(None, search_query)
            if status != "OK":
                raise IntegrationError(f"IMAP search failed: {status}")

            email_ids = messages[0].split()[-limit:] if limit > 0 else messages[0].split()
            parsed_emails: list[ParsedEmail] = []

            for email_id in email_ids:
                status, msg_data = self._imap_connection.fetch(email_id, "(RFC822)")
                if status != "OK":
                    logger.warning(f"Failed to fetch email {email_id}")
                    continue

                for response_part in msg_data:
                    if isinstance(response_part, tuple):
                        email_bytes = response_part[1]
                        parsed_email = await self.parse_email_bytes(email_bytes)
                        if parsed_email:
                            parsed_emails.append(parsed_email)

            return parsed_emails
        except Exception as e:
            raise IntegrationError(f"Failed to fetch emails: {e}") from e

    async def parse_email_file(self, file_path: str) -> ParsedEmail | None:
        try:
            with open(file_path, "rb") as f:
                email_bytes = f.read()
            return await self.parse_email_bytes(email_bytes)
        except Exception as e:
            logger.error(f"Failed to parse email file {file_path}: {e}")
            return None

    async def parse_email_bytes(self, email_bytes: bytes) -> ParsedEmail | None:
        if not MAILPARSER_AVAILABLE:
            logger.error("mailparser library is not available")
            raise IntegrationError("mailparser library is required for email parsing")

        try:
            mail = mailparser.parse_from_bytes(email_bytes)

            message_id = mail.message_id or str(uuid4())
            subject = mail.subject
            sender = mail.from_[0][1] if mail.from_ else None
            received_at = mail.date or datetime.now()

            attachments: list[ParsedAttachment] = []
            for attachment in mail.attachments:
                parsed_attachment = await self._parse_attachment(
                    attachment, message_id, received_at
                )
                if parsed_attachment:
                    if parsed_attachment.hash in self._processed_hashes:
                        logger.info(f"Skipping duplicate attachment: {parsed_attachment.file_name}")
                        continue
                    self._processed_hashes.add(parsed_attachment.hash)
                    attachments.append(parsed_attachment)

            return ParsedEmail(
                message_id=message_id,
                subject=subject,
                sender=sender,
                received_at=received_at,
                body=mail.body,
                attachments=attachments,
            )
        except Exception as e:
            logger.error(f"Failed to parse email: {e}")
            return None

    async def _parse_attachment(
        self,
        attachment: dict[str, Any],
        message_id: str,
        received_at: datetime,
    ) -> ParsedAttachment | None:
        try:
            file_name = attachment.get("filename", f"attachment_{uuid4()}")
            content_type = attachment.get("mail_content_type", "")
            payload = attachment.get("payload", b"")

            if isinstance(payload, str):
                payload = payload.encode("utf-8")

            file_hash = hashlib.sha256(payload).hexdigest()
            file_size = len(payload)

            config = self.config
            if not isinstance(config, EmailConfig):
                config = EmailConfig(**config.model_dump())

            os.makedirs(config.attachment_dir, exist_ok=True)
            file_path = os.path.join(
                config.attachment_dir,
                f"{message_id}_{file_name}",
            )

            with open(file_path, "wb") as f:
                f.write(payload)

            attachment_type = await self._detect_attachment_type(file_name, content_type)
            content = await self._extract_attachment_content(file_path, content_type)

            case_no = await self._extract_case_no(content or "", file_name)
            invoice_no = await self._extract_invoice_no(content or "", file_name)
            amount = await self._extract_amount(content or "")

            return ParsedAttachment(
                file_name=file_name,
                file_path=file_path,
                file_size=file_size,
                content_type=content_type,
                hash=file_hash,
                attachment_type=attachment_type,
                content=content,
                case_no=case_no,
                invoice_no=invoice_no,
                amount=amount,
            )
        except Exception as e:
            logger.error(f"Failed to parse attachment: {e}")
            return None

    async def _detect_attachment_type(
        self,
        file_name: str,
        content_type: str,
    ) -> AttachmentType:
        file_name_lower = file_name.lower()
        content_type_lower = content_type.lower()

        if any(kw in file_name_lower for kw in ["invoice", "发票", "账单"]):
            return AttachmentType.INVOICE
        elif any(kw in file_name_lower for kw in ["bill", "费用"]):
            return AttachmentType.BILL
        elif any(kw in file_name_lower for kw in ["contract", "协议", "合同"]):
            return AttachmentType.CONTRACT
        elif any(kw in file_name_lower for kw in ["receipt", "收据"]):
            return AttachmentType.RECEIPT
        elif any(kw in file_name_lower for kw in ["statement", "对账单"]):
            return AttachmentType.STATEMENT

        if "pdf" in content_type_lower or file_name_lower.endswith(".pdf"):
            return AttachmentType.INVOICE

        return AttachmentType.OTHER

    async def _extract_attachment_content(
        self,
        file_path: str,
        content_type: str,
    ) -> str | None:
        content_type_lower = content_type.lower()
        file_name_lower = file_path.lower()

        try:
            if "pdf" in content_type_lower or file_name_lower.endswith(".pdf"):
                return await self._extract_pdf_content(file_path)
            elif (
                "excel" in content_type_lower
                or "spreadsheet" in content_type_lower
                or file_name_lower.endswith((".xlsx", ".xls"))
            ):
                return await self._extract_excel_content(file_path)
            elif (
                "word" in content_type_lower
                or file_name_lower.endswith((".docx", ".doc"))
            ):
                return await self._extract_word_content(file_path)
            elif "text" in content_type_lower or file_name_lower.endswith(".txt"):
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
        except Exception as e:
            logger.warning(f"Failed to extract content from {file_path}: {e}")

        return None

    async def _extract_pdf_content(self, file_path: str) -> str | None:
        try:
            from PyPDF2 import PdfReader

            reader = PdfReader(file_path)
            text_parts: list[str] = []
            for page in reader.pages:
                text_parts.append(page.extract_text() or "")
            return "\n".join(text_parts)
        except ImportError:
            logger.warning("PyPDF2 not installed, cannot extract PDF content")
            return None
        except Exception as e:
            logger.warning(f"Failed to extract PDF content: {e}")
            return None

    async def _extract_excel_content(self, file_path: str) -> str | None:
        if not OPENPYXL_AVAILABLE:
            logger.warning("openpyxl library is not available, cannot extract Excel content")
            return None

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text_parts: list[str] = []

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text_parts.append(f"--- Sheet: {sheet_name} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_str = " | ".join(str(cell) for cell in row if cell is not None)
                    if row_str:
                        text_parts.append(row_str)

            return "\n".join(text_parts)
        except Exception as e:
            logger.warning(f"Failed to extract Excel content: {e}")
            return None

    async def _extract_word_content(self, file_path: str) -> str | None:
        try:
            from docx import Document

            doc = Document(file_path)
            text_parts: list[str] = []

            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)

            for table in doc.tables:
                for row in table.rows:
                    row_str = " | ".join(cell.text for cell in row.cells if cell.text.strip())
                    if row_str:
                        text_parts.append(row_str)

            return "\n".join(text_parts)
        except ImportError:
            logger.warning("python-docx not installed, cannot extract Word content")
            return None
        except Exception as e:
            logger.warning(f"Failed to extract Word content: {e}")
            return None

    async def _extract_case_no(self, content: str, file_name: str) -> str | None:
        patterns = [
            r'[(（](\d{4})[)）][\u4e00-\u9fa5]*?民初字第(\d+)号',
            r'[(（](\d{4})[)）][\u4e00-\u9fa5]*?民初字第(\d+)号',
            r'CASE[-_]?(\d+)',
            r'案号[:：]?\s*(\S+)',
        ]

        for pattern in patterns:
            match = re.search(pattern, content)
            if match:
                return match.group(0)

        match = re.search(r'[(（](\d{4})[)）](\d+)号', file_name)
        if match:
            return match.group(0)

        return None

    async def _extract_invoice_no(self, content: str, file_name: str) -> str | None:
        patterns = [
            r'发票号码[:：]?\s*(\d{8,20})',
            r'Invoice\s*(?:No|Number)[:：]?\s*([A-Z0-9-]+)',
            r'INV[-_]?(\d+)',
            r'账单编号[:：]?\s*(\S+)',
        ]

        for pattern in patterns:
            match = re.search(pattern, content, re.IGNORECASE)
            if match:
                return match.group(1)

        match = re.search(r'INV[-_]?([A-Z0-9-]+)', file_name, re.IGNORECASE)
        if match:
            return match.group(1)

        return None

    async def _extract_amount(self, content: str) -> float | None:
        patterns = [
            r'金额[:：]?\s*(?:¥|RMB|CNY)?\s*([\d,]+\.?\d*)',
            r'合计[:：]?\s*(?:¥|RMB|CNY)?\s*([\d,]+\.?\d*)',
            r'Total[:：]?\s*(?:¥|RMB|CNY)?\s*([\d,]+\.?\d*)',
            r'Amount[:：]?\s*(?:¥|RMB|CNY)?\s*([\d,]+\.?\d*)',
            r'(?:¥|RMB|CNY)\s*([\d,]+\.?\d*)',
        ]

        for pattern in patterns:
            match = re.search(pattern, content, re.IGNORECASE)
            if match:
                try:
                    amount_str = match.group(1).replace(",", "")
                    return float(amount_str)
                except (ValueError, TypeError):
                    continue

        return None

    async def transform_data(self, raw_data: list[Any]) -> list[ParsedEmail]:
        return [
            data if isinstance(data, ParsedEmail) else ParsedEmail(**data)
            for data in raw_data
        ]

    async def link_to_invoice(
        self,
        parsed_email: ParsedEmail,
    ) -> dict[str, Any]:
        if not self.db:
            return {"linked": False, "reason": "No database session"}

        results: list[dict[str, Any]] = []

        for attachment in parsed_email.attachments:
            if not attachment.invoice_no or not attachment.case_no:
                results.append({
                    "file_name": attachment.file_name,
                    "linked": False,
                    "reason": "Missing invoice_no or case_no",
                })
                continue

            try:
                case_result = await self.db.execute(
                    select(Invoice).where(
                        Invoice.invoice_no == attachment.invoice_no
                    )
                )
                invoice = case_result.scalar_one_or_none()

                if invoice:
                    email_attachment = EmailAttachment(
                        message_id=parsed_email.message_id,
                        subject=parsed_email.subject,
                        sender=parsed_email.sender,
                        received_at=parsed_email.received_at,
                        file_name=attachment.file_name,
                        file_path=attachment.file_path,
                        hash=attachment.hash,
                        linked_invoice_id=invoice.id,
                    )
                    self.db.add(email_attachment)
                    await self.db.commit()

                    results.append({
                        "file_name": attachment.file_name,
                        "linked": True,
                        "invoice_id": str(invoice.id),
                        "invoice_no": attachment.invoice_no,
                    })
                else:
                    if attachment.amount and attachment.invoice_no:
                        invoice_create = await self._create_invoice_from_attachment(
                            attachment
                        )
                        results.append({
                            "file_name": attachment.file_name,
                            "linked": False,
                            "reason": "Invoice not found",
                            "suggested_invoice": invoice_create.model_dump() if invoice_create else None,
                        })
                    else:
                        results.append({
                            "file_name": attachment.file_name,
                            "linked": False,
                            "reason": "Invoice not found",
                        })
            except Exception as e:
                logger.error(f"Failed to link attachment {attachment.file_name}: {e}")
                results.append({
                    "file_name": attachment.file_name,
                    "linked": False,
                    "error": str(e),
                })

        return {
            "message_id": parsed_email.message_id,
            "total_attachments": len(parsed_email.attachments),
            "results": results,
        }

    async def _create_invoice_from_attachment(
        self,
        attachment: ParsedAttachment,
    ) -> InvoiceCreate | None:
        if not attachment.amount or not attachment.invoice_no:
            return None

        try:
            return InvoiceCreate(
                invoice_no=attachment.invoice_no,
                case_id=UUID(int=0),
                amount=attachment.amount,
                status=InvoiceStatus.PENDING,
                invoice_date=datetime.now().date(),
                source=InvoiceSource.EMAIL,
                items=[
                    InvoiceItemCreate(
                        item_name=attachment.file_name,
                        description=attachment.attachment_type.value,
                        quantity=1,
                        unit_price=attachment.amount,
                        amount=attachment.amount,
                        fee_type=attachment.attachment_type.value,
                    )
                ],
            )
        except Exception as e:
            logger.error(f"Failed to create invoice from attachment: {e}")
            return None

    async def _generate_mock_data(
        self,
        limit: int = 50,
        *args: Any,
        **kwargs: Any,
    ) -> list[ParsedEmail]:
        senders = [
            "billing@lawfirm.com",
            "invoice@court.gov.cn",
            "finance@client.com",
            "contract@partner.com",
            "admin@legal.org",
        ]

        subjects = [
            "发票 - {case_no} 法律服务费",
            "账单通知 - {case_no}",
            "合同附件 - {case_name}",
            "付款提醒 - Invoice {invoice_no}",
            "案件进展通知 - {case_no}",
        ]

        case_names = [
            "张某合同纠纷案",
            "王某知识产权案",
            "刘某劳动争议案",
            "陈某股权转让案",
        ]

        attachment_types = [
            (AttachmentType.INVOICE, "invoice_{invoice_no}.pdf", "application/pdf"),
            (AttachmentType.BILL, "bill_{case_no}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            (AttachmentType.CONTRACT, "contract_{case_no}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            (AttachmentType.RECEIPT, "receipt_{invoice_no}.pdf", "application/pdf"),
        ]

        emails: list[ParsedEmail] = []
        count = min(limit, 20)

        for i in range(count):
            year = datetime.now().year
            case_no = f"({year})京民初字第{i + 1:06d}号"
            invoice_no = f"INV-{year}-{i + 1:04d}"
            case_name = random.choice(case_names)

            attachment_type, file_name_pattern, content_type = random.choice(attachment_types)
            file_name = file_name_pattern.format(invoice_no=invoice_no, case_no=case_no)

            amount = round(random.uniform(5000, 100000), 2)
            content = (
                f"发票号码: {invoice_no}\n"
                f"案号: {case_no}\n"
                f"案件名称: {case_name}\n"
                f"金额: ¥{amount}\n"
                f"合计: ¥{amount}\n"
            )

            file_hash = hashlib.sha256(file_name.encode()).hexdigest()

            attachment = ParsedAttachment(
                file_name=file_name,
                file_path=f"/tmp/{file_name}",
                file_size=random.randint(10000, 500000),
                content_type=content_type,
                hash=file_hash,
                attachment_type=attachment_type,
                content=content,
                case_no=case_no,
                invoice_no=invoice_no,
                amount=amount,
            )

            subject = random.choice(subjects).format(
                case_no=case_no,
                invoice_no=invoice_no,
                case_name=case_name,
            )

            email = ParsedEmail(
                message_id=f"<mock-{i}@example.com>",
                subject=subject,
                sender=random.choice(senders),
                received_at=datetime.now() - timedelta(days=random.randint(0, 30)),
                body=f"请查收附件：{subject}\n\n此邮件由系统自动发送，请勿直接回复。",
                attachments=[attachment],
            )
            emails.append(email)

        logger.info(f"Generated {len(emails)} mock emails")
        return emails
